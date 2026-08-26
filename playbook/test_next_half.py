# -*- coding: utf-8 -*-
"""다음 반기 데이터로도 돌아가는지 시험한다.

상반기 파일을 +6개월 밀어 '하반기'인 척 만들고 run.py를 그대로 태운다.
반기가 바뀌면 없어지는 지표(예: 2월)나 시즌 캘린더 연도 같은 사고를 여기서 잡는다.

  python3 -m playbook.test_next_half NOM.xlsx SOM.xlsx COM.xlsx --reject stats.xlsx
"""
import re
import sys
from pathlib import Path

import pandas as pd
from playbook import ingest

_orig = ingest.load_all


def shifted(paths):
    dfs = _orig(paths)
    for d in dfs.values():
        d['일시'] = d['일시'] + pd.DateOffset(months=6)
        d['월'] = d['일시'].dt.month
        d['요일'] = d['일시'].dt.dayofweek
        d['시'] = d['일시'].dt.hour
    return dfs


ingest.load_all = shifted

from playbook import run  # noqa: E402  (monkeypatch 먼저)

out = Path(__file__).resolve().parent / '_test_하반기.pptx'
sys.argv = ['run.py'] + sys.argv[1:] + ['--out', str(out)]
print('=== 하반기 시뮬레이션 ===')
code = run.main()

# 결과 점검
from pptx import Presentation  # noqa: E402
prs = Presentation(str(out))
bad, cal_headings = [], []
for si, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if '상반기' in t and '시즌 캘린더' not in t:
            bad.append((si, t[:60]))
        # 캘린더 표제는 카드가 실제로 덮는 달을 가리켜야 한다.
        # 하반기(7~12월) 데이터면 2027년 3~6월을 안내한다(데이터 끝 +3개월부터 4개월).
        if re.match(r'\d{4}\s*\d{1,2}~\d{1,2}월\s*(시즌\s*)?캘린더', t.strip()):
            cal_headings.append((si, t.strip()))
        elif re.search(r'\d{4}\s*(상반기|하반기)\s*(시즌\s*)?캘린더', t):
            bad.append((si, f'캘린더 표제가 안 바뀜: {t[:60]}'))
for si, t in cal_headings:
    if '2027 3~6월' not in t:
        bad.append((si, f'캘린더가 2027 3~6월을 안 가리킴: {t}'))
if not cal_headings:
    bad.append((0, '캘린더 표제를 하나도 못 찾았습니다'))

print()
if bad:
    print(f'문제 {len(bad)}건')
    for si, t in bad:
        print(f'  slide {si}: {t}')
else:
    print(f'기간 표기 점검 통과 — 남은 "상반기" 없음, '
          f'캘린더 {len(cal_headings)}곳 모두 2027 3~6월')
sys.exit(code or (1 if bad else 0))
