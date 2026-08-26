# -*- coding: utf-8 -*-
"""슬라이드의 도형ID와 텍스트를 보여준다. overrides.py에 자리를 적을 때 쓴다.

  python shape_id.py 36
"""
import sys
from pptx import Presentation
from pptx.util import Emu

prs = Presentation(sys.argv[2] if len(sys.argv) > 2 else 'template_v2.pptx')
for n in [int(x) for x in sys.argv[1].split(',')]:
    print(f'=== slide {n}')
    for sh in sorted(prs.slides[n - 1].shapes, key=lambda s: (Emu(s.top).inches, Emu(s.left).inches)):
        t = sh.text_frame.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        if t:
            print(f'  id {sh.shape_id:4d}  {t[:95]}')
